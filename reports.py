"""Excel risk pack."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def build_xlsx(s, name):
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    def sheet(title, head, rows, fmt=None):
        w = ws if title == "Summary" else wb.create_sheet(title)
        w.append(head)
        for c in w[w.max_row]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="14212E")
        for r in rows: w.append(list(r))
        for i, col in enumerate(w.columns):
            w.column_dimensions[col[0].column_letter].width = 22 if i else 30
            for c in col[1:]:
                if fmt and i in fmt and isinstance(c.value, (int, float)): c.number_format = fmt[i]
    sheet("Summary", ["Item", "Value"], [("Portfolio", name), ("As of", s["as_of"]), ("Gross exposure (INR)", s["nav"]), ("Position lines", s["n_lines"]), ("Instruments", s["n_instruments"]),
        ("Data source", s["data_source"]), ("Backtest breaches (250d)", s["backtest"]["breaches"]), ("Backtest zone", s["backtest"]["zone"]), ("Limit warning / critical", f"{s['limits']['warn']:.2%} / {s['limits']['crit']:.2%}")], {1: "#,##0"})
    sheet("Risk", ["Method", "VaR 99% (INR)", "ES 99% (INR)"], [(k.replace("_", " "), d["var"], d["es"]) for k, d in s["methods"].items()], {1: "#,##0", 2: "#,##0"})
    sheet("Attribution", ["Level", "Key", "Component VaR (INR)", "% of risk"], [(lv, r["key"], r["var"], r["pct"]) for lv, rs in s["attribution"].items() for r in rs], {2: "#,##0", 3: "0.00%"})
    sheet("Stress", ["Scenario", "P&L (INR)", "% of exposure", "Largest losses"], [(x["name"], x["loss"], x["pct"], ", ".join(f"{t['symbol']} {t['pnl']:,.0f}" for t in x["top"])) for x in s["stress"]], {1: "#,##0", 2: "0.00%"})
    sheet("Alerts", ["Severity", "Title", "Explanation", "As of"], [(a["severity"], a["title"], a["explanation"], a["as_of"]) for a in s["alerts"]])
    sheet("Positions", ["Symbol", "Name", "Class", "Ccy", "Sector", "Value (INR)", "Weight", "Component VaR", "% of risk"],
          [(p["symbol"], p["name"], p["cls"], p["ccy"], p["sector"], p["value"], p["weight"], p["comp_var"], p["pct"]) for p in s["positions"]], {5: "#,##0", 6: "0.00%", 7: "#,##0", 8: "0.00%"})
    b = io.BytesIO(); wb.save(b); return b.getvalue()
