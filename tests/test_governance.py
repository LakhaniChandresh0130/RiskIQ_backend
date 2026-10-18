import engine as E, repo, reports, notify, llm, json, io
from db import DB
from service import Service, NotFound

def fresh(): db = DB(":memory:"); return db, Service(db)
def raises(exc, fn, *a):
    try: fn(*a)
    except exc: return True
    return False

def test_first_login_creates_org_and_invite_joins_it():
    db, _ = fresh(); a = repo.resolve_user(db, "u1", "a@x.com"); assert a["role"] == "admin"
    assert repo.resolve_user(db, "u1", "a@x.com")["org"] == a["org"]
    repo.invite(db, a["org"], "b@x.com", "viewer", "u1"); b = repo.resolve_user(db, "u2", "b@x.com")
    assert b["org"] == a["org"] and b["role"] == "viewer"
    assert raises(ValueError, repo.invite, db, a["org"], "b@x.com", "viewer", "u1")

def test_role_permissions_and_last_admin_protected():
    assert not repo.can("viewer", "write") and repo.can("analyst", "write") and not repo.can("analyst", "approve") and repo.can("auditor", "audit") and not repo.can("risk_manager", "admin")
    db, _ = fresh(); a = repo.resolve_user(db, "u1", "a@x.com"); mid = repo.list_members(db, a["org"])[0]["id"]
    assert raises(ValueError, repo.set_role, db, a["org"], mid, "viewer", "u1")

def test_portfolios_are_isolated_between_organisations():
    db, svc = fresh(); a, b = repo.resolve_user(db, "u1", "a@x.com"), repo.resolve_user(db, "u2", "b@x.com")
    pid = repo.create_portfolio(db, a["org"], "u1", "P", E.demo_large(300, 40), "test")
    assert repo.get_portfolio(db, b["org"], pid) is None and raises(NotFound, svc.analysis, b["org"], pid)
    assert svc.snapshot(a["org"], pid)["n_lines"] == 300

def test_limits_maker_checker_and_effect_on_alerts():
    db, svc = fresh(); a = repo.resolve_user(db, "u1", "a@x.com"); org = a["org"]
    assert raises(ValueError, repo.propose_limits, db, org, "demo", .05, .03, "u1")
    lid = repo.propose_limits(db, org, "demo", .0005, .001, "u1")
    assert raises(PermissionError, repo.approve_limits, db, org, lid, "u1")          # proposer cannot approve
    assert repo.active_limits(db, org, "demo", E.LIMITS) == E.LIMITS               # nothing changes until approved
    repo.approve_limits(db, org, lid, "u2"); s = svc.snapshot(org, "demo")
    assert s["limits"]["crit"] == .001 and any(x["rule"] == "var_limit" and x["severity"] == "critical" for x in s["alerts"])

def test_scenarios_validated_approved_and_used():
    db, svc = fresh(); org = repo.resolve_user(db, "u1", "a@x.com")["org"]
    assert raises(ValueError, repo.propose_scenario, db, org, "bad", {"foo": 1}, "u1") and raises(ValueError, repo.propose_scenario, db, org, "big", {"eq_in": 5}, "u1")
    sid = repo.propose_scenario(db, org, "RBI shock", {"eq_in": -.1, "dy": .01}, "u1")
    assert "RBI shock" not in [x["name"] for x in svc.snapshot(org, "demo")["stress"]]
    repo.approve_scenario(db, org, sid, "u2"); assert "RBI shock" in [x["name"] for x in svc.snapshot(org, "demo")["stress"]]

def test_alerts_are_stored_once_notified_once_and_lifecycle_works():
    db, svc = fresh(); org = repo.resolve_user(db, "u1", "a@x.com")["org"]; sent = []; notify.deliver = lambda s, t: sent.append(s)
    snap = svc.snapshot(org, "demo"); n = len(snap["alerts"]); assert n > 0 and len(sent) == n
    svc.cache.clear(); svc.snapshot(org, "demo"); assert len(sent) == n                 # same alerts are not re-sent
    aid = snap["alerts"][0]["id"]; repo.act_alert(db, org, aid, "u1", "acknowledged")
    assert len(svc.snapshot(org, "demo")["alerts"]) == n - 1 and len(repo.list_alerts(db, org, "demo", ("acknowledged",))) == 1
    assert raises(ValueError, repo.act_alert, db, org, "nope", "u1", "resolved")

def test_audit_trail_records_actions():
    db, svc = fresh(); org = repo.resolve_user(db, "u1", "a@x.com")["org"]
    repo.create_portfolio(db, org, "u1", "P", E.demo_large(200, 30), "t"); repo.propose_limits(db, org, "demo", .01, .02, "u1")
    acts = {r["action"] for r in repo.list_audit(db, org)}; assert {"org.create", "portfolio.create", "limits.propose"} <= acts

def test_excel_report_and_scheduled_refresh():
    db, svc = fresh(); org = repo.resolve_user(db, "u1", "a@x.com")["org"]; notify.deliver = lambda s, t: None
    pid = repo.create_portfolio(db, org, "u1", "P", E.demo_large(500, 60), "t"); s = svc.snapshot(org, pid)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(reports.build_xlsx(s, "P"))); assert wb.sheetnames == ["Summary", "Risk", "Attribution", "Stress", "Alerts", "Positions"] and wb["Positions"].max_row > 5
    assert svc.refresh_all() == 1

def test_groq_tool_loop_uses_openai_format_and_returns_grounded_answer():
    seen = []
    def fake(url, payload, key):
        seen.append(payload)
        if len(seen) == 1: return {"choices": [{"message": {"content": None, "tool_calls": [{"id": "c1", "type": "function", "function": {"name": "get_alerts", "arguments": "{}"}}]}}]}
        return {"choices": [{"message": {"content": "VaR is elevated; see alert."}}]}
    import os; os.environ["GROQ_API_KEY"] = "gsk_test"
    tools = [{"name": "get_alerts", "description": "d", "input_schema": {"type": "object", "properties": {}}}]
    ans, used = llm.groq_chat("why?", tools, "sys", lambda n, a: [{"title": "x"}], post=fake)
    assert used == ["get_alerts"] and ans.startswith("VaR is elevated") and seen[0]["tools"][0]["function"]["name"] == "get_alerts"
    assert seen[1]["messages"][-1]["role"] == "tool" and llm.provider() == "groq"
    os.environ.pop("GROQ_API_KEY")

def test_provider_priority_and_failure_falls_back():
    import os; os.environ.pop("GROQ_API_KEY", None); os.environ.pop("ANTHROPIC_API_KEY", None)
    assert llm.provider() is None and llm.ask("q", "s", [], lambda n, a: {}) is None
    os.environ["GROQ_API_KEY"] = "bad"; llm._post = lambda *a: (_ for _ in ()).throw(RuntimeError("429"))
    llm.groq_chat.__defaults__ = (llm._post,); assert llm.ask("q", "s", [], lambda n, a: {}) is None; os.environ.pop("GROQ_API_KEY")

def test_zerodha_style_export_columns_are_recognised():
    import ingest
    csv = (b"Symbol,ISIN,Sector,Quantity Available,Average Price,Previous Closing Price\n"
           b"RELIANCE,INE002A01018,Energy,100,2400,2900\nHDFCBANK,INE040A01034,Financials,50,1500,1650\nTCS,INE467B01029,IT,10,3000,3500\n")
    rows, rep = ingest.parse(csv, "holdings.csv")
    assert rep["rows_ok"] == 3 and rows[0]["value"] == 290000 and rows[0]["sector"] == "Energy" and rows[0]["yf"] == "RELIANCE.NS"
    s, _ = E.analyze(rows); assert s["n_lines"] == 3 and s["methods"]["parametric"]["var"] > 0
def test_unknown_columns_give_helpful_error():
    import ingest
    try: ingest.parse(b"foo,bar\n1,2\n", "x.csv")
    except ValueError as e: assert "Columns found: foo, bar" in str(e)
    else: assert False
